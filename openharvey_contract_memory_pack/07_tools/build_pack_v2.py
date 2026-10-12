#!/usr/bin/env python3
"""Build business-facing documents and operator-only source locations. No model calls."""
import hashlib,json,re
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
NUMBERS={x:f'HZ-{group}-2026-{i:03d}' for group,items in [('CG',['H01','H03','H04','H06','F01','F03','F04']),('XS',['H02','H05','F02'])] for i,x in enumerate(items,1)}
PROJECTS=dict(zip(['PRJ-PLANT-A','PRJ-QH-PARTS','PRJ-YL-PILOT','PRJ-CATALOG-A','PRJ-JS-OPS','PRJ-EVENT-APR','PRJ-PLANT-B','PRJ-YL-CONNECT','PRJ-YL-ANNUAL','PRJ-CATALOG-B'],['一期生产线终端采购项目','终端配套组件供货项目','流程协作软件六个月试点项目','商品目录清洗项目','设备运行数据委托处理项目','四月产品交流会项目','二期生产线终端采购项目','协同接口开发项目','流程协作软件年度项目','商品知识库加工项目']))
ROLES={'ROLE-CFO':'财务负责人','ROLE-LEGAL-LEAD':'法务负责人','ROLE-BUSINESS':'业务经办人','ROLE-BUSINESS-OWNER':'业务负责人','ROLE-AGENT':'审查助手','ROLE-ASSISTANT':'审查助手','ROLE-PROCUREMENT':'采购负责人','ROLE-IT':'信息技术负责人','ROLE-BOTH-PARTIES-AUTHORIZED':'双方授权代表','ROLE-SECURITY-REVIEWER':'安全审核负责人','ROLE-FINANCE-VERIFIER':'财务核验负责人','ROLE-DATA-SECURITY':'数据安全负责人'}
REFERENCES={}
def dump(path,data):
    (ROOT/path).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
def rows(path):return [json.loads(l) for l in (ROOT/path).read_text().splitlines() if l]
def business(s):
    for a,b in sorted(REFERENCES.items(),key=lambda x:-len(x[0])):s=s.replace(a,b)
    for a,b in PROJECTS.items():s=s.replace(a,b)
    for a,b in sorted(NUMBERS.items(),key=lambda x:-len(x[0])):s=re.sub(r'(?<![A-Za-z0-9])'+a+r'(?![A-Za-z0-9])',b,s)
    for a,b in sorted(ROLES.items(),key=lambda x:-len(x[0])):s=s.replace(a,b)
    s=s.replace('本测试中','本公司').replace('本测试','本公司').replace('本样本','本合同')
    s=s.replace('虚构双方补充约定摘要：','双方补充约定摘要：').replace('虚构执行证据：','履行核验记录：')
    s=s.replace('本公司视该补充约定已确认，仅作为合成执行事实，不附真实签署件。','本记录留存补充约定的处理结论，签署文件另行归档。')
    s=s.replace('该记录仅为测试事实，不附真实银行文件，不代表续约或其他合同满足条件。','银行保函核验资料由财务另行归档，本记录不代表续约或其他合同满足条件。')
    s=s.replace('合成事实中','业务记录中').replace('不代表现实授权或其他合同','适用范围限于本补充约定')
    s=s.replace('合成测试记录','脱敏的流程验证记录')
    return s
def docx(path,body):
    d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.1)
    for name in ['Normal','Title','Heading 1']:
        st=d.styles[name];st.font.name='Arial Unicode MS';fonts=st._element.get_or_add_rPr().rFonts;fonts.attrib.clear()
        for script in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+script),'Arial Unicode MS')
        st.font.color.rgb=RGBColor(0,0,0)
    d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.2
    d.styles['Normal'].paragraph_format.space_after=Pt(5)
    d.styles['Title'].font.size=Pt(19);d.styles['Heading 1'].font.size=Pt(12);d.styles['Heading 1'].font.bold=False
    lines=body.splitlines();d.add_paragraph(lines[0],'Title')
    for line in lines[1:]:
        if not line:continue
        if re.match(r'第\d+条 ',line):p=d.add_paragraph(line,'Heading 1');p.paragraph_format.keep_with_next=True
        else:p=d.add_paragraph(line)
    for st in d.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.core_properties.title=lines[0];d.core_properties.author='';d.core_properties.comments=''
    d.save(ROOT/path)
def main():
    manifest=json.loads((ROOT/'00_start_here/contract_manifest.json').read_text())
    manifest=[m for m in manifest if m['doc_id']!='F01-R2']
    mappings=[]
    for m in manifest:
        cid=m['doc_id'];body=(ROOT/m['txt_path']).read_text()
        body='\n'.join(l for l in body.splitlines() if not any(w in l for w in ['测试主体映射','本文件为原创虚构','来源标识：']))
        body=re.sub(r' \[(?:H|F)\d{2}-C\d{2}\]','',body)
        body=body.replace('双方主体映射见首页测试元数据。','双方应在签署前核验主体名称和授权文件。')
        body=body.replace('为防止误用，本合同不提供真实地址、自然人、电话、印章或银行账号；正式签署前应另附经核验的主体、送达和账户清单。','双方的联系信息、送达地址及收款账户以经核验的书面清单为准，签署前由各方有权代表确认。')
        body=body.replace('为防止误用，本样本不提供真实地址、自然人、电话、印章或银行账号；正式签署前应另附经核验的主体、送达和账户清单。','双方的联系信息、送达地址及收款账户以经核验的书面清单为准，签署前由各方有权代表确认。')
        body=body.replace('；本测试文件未签署、未盖章，不代表现实法律承诺','')
        body=body.replace('甲方签署栏与乙方签署栏均留空；签署日期以真实签署时填写为准。','')
        body=business(body).strip()
        if '甲方签章：' not in body:body+='\n\n甲方签章：________________    乙方签章：________________\n授权代表：________________    授权代表：________________\n签署日期：________________    签署日期：________________'
        body+='\n';(ROOT/m['txt_path']).write_text(body);docx(m['docx_path'],body)
        m.update(contract_number=NUMBERS[cid],version='R1',text_characters=len(body))
    original=next(m for m in manifest if m['doc_id']=='F01')
    revision=dict(original,doc_id='F01-R2',version='R2',txt_path='03_future/contracts_txt/F01_R2_工业控制终端补充采购合同.txt',docx_path='03_future/contracts_docx/F01_R2_工业控制终端补充采购合同.docx')
    body=(ROOT/original['txt_path']).read_text().replace('草案日期 2026-08-03','修订日期 2026-08-21  |  修订版本 R2')
    parts=re.split(r'(?=第\d+条 )',body);parts=[p.replace('0.03%','0.05%') if p.startswith('第6条 ') else p for p in parts];body=''.join(parts)
    (ROOT/revision['txt_path']).write_text(body);docx(revision['docx_path'],body);revision['text_characters']=len(body);manifest.append(revision)
    for m in manifest:
        body=(ROOT/m['txt_path']).read_text();digest=hashlib.sha256(body.encode()).hexdigest()
        m['txt_sha256']=digest;m['docx_sha256']=hashlib.sha256((ROOT/m['docx_path']).read_bytes()).hexdigest()
        clauses=list(re.finditer(r'^第(\d+)条 ([^\n]+)\n',body,re.M))
        for i,c in enumerate(clauses):
            end=clauses[i+1].start() if i+1<len(clauses) else body.find('\n甲方签章：',c.end())
            text=body[c.start():end].strip();anchor=f"{m['doc_id']}-C{int(c.group(1)):02d}"
            mappings.append({'anchor':anchor,'input_id':m['doc_id'],'file':m['txt_path'],'file_sha256':digest,'contract_number':m['contract_number'],'version':m['version'],'clause_number':int(c.group(1)),'clause_title':c.group(2),'text':text,'text_sha256':hashlib.sha256(text.encode()).hexdigest()})
    events=[]
    for n in ['reviewer_comments','final_decisions_and_updates','policies']:events+=rows('02_review_writeback/operator_only_indexes/'+n+'.jsonl')
    for i,e in enumerate(sorted(events,key=lambda e:(e['recorded_at'],e['event_id'])),1):
        for k in ['event_id','source_doc_id','source_clause_id']:REFERENCES[e[k]]=f'HZ-JL-2026-{i:03d}'
    for i,e in enumerate(sorted(events,key=lambda e:(e['recorded_at'],e['event_id'])),1):
        kind={'human_opinion':'审阅意见记录','agent_suggestion':'审查建议记录','execution_result':'履行情况记录'}.get(e['kind'],'处理决定记录')
        if e['event_id'].startswith('POL-'):kind='内部管理规定'
        if e['event_id'].startswith('WD-'):kind='内部先例撤回通知'
        if e['event_id'].startswith('SUP-'):kind='补充约定记录'
        authority=business(e['authority_scope']);content=business(e['content'])
        role=business(e['actor_role_id']).replace('+','、');role=re.sub(r'ROLE-[A-Z-]+','业务负责人',role)
        body=f'{kind}\n编号：HZ-JL-2026-{i:03d}\n记录时间：{e["recorded_at"]}\n适用时间：{e["effective_at"]}\n记录角色：{role}\n职责范围：{authority}\n\n{content}\n'
        path=f'02_review_writeback/event_files/{e["recorded_at"][:10]}_{e["event_id"]}.txt';(ROOT/path).write_text(body)
        mappings.append({'anchor':e['source_clause_id'],'input_id':e['event_id'],'file':path,'file_sha256':hashlib.sha256(body.encode()).hexdigest(),'clause_number':None,'clause_title':kind,'text':content,'text_sha256':hashlib.sha256(content.encode()).hexdigest()})
    identity='合同主体名称及简称对照\n\n华子有限公司，简称华子，为本公司主体。\n启衡设备有限公司，简称启衡设备或启衡。\n云砺信息技术有限公司，简称云砺信息或云砺。\n景澍数字服务有限公司，简称景澍数字或景澍。\n棠序会展服务有限公司，简称棠序会展或棠序。\n\n简称仅用于文书阅读，主体判断应以完整名称及核验资料为准。近似名称不自动视为同一主体；甲方和乙方表示本份合同中的角色，不是固定公司身份。\n'
    (ROOT/'01_history/identity_map.txt').write_text(identity)
    (ROOT/'06_negative_controls_QUARANTINE/外部软件采购审批记录.txt').write_text('软件采购例外审批记录\n编号：HZ-WB-2026-001\n记录日期：2026年8月1日\n采购主体：华子有限公司\n供应商：云砺信息技术有限公司\n\n财务和法务负责人同意本公司本项目采用90%预付款，无须提供担保。此意见仅适用于本公司本项目，其他公司的合同不属于本次审批范围。\n')
    for p in [ROOT/'04_queries/queries.jsonl',ROOT/'04_queries/queries.txt']:
        if p.suffix=='.jsonl':
            q=rows(str(p.relative_to(ROOT)))
            for x in q:
                x['prompt']=business(x['prompt'])
                if x['query_id']=='Q21':x.update(prompt='请依据当前可用的记录审阅这份试点合同的付款和责任安排。区分已有处理记录与已经确认可复用的经验，不推定条件已满足。',as_of='2026-02-09T09:00:00+08:00')
                if x['query_id']=='Q22':x.update(prompt='刚加入了撤回记录。请核对以前关于训练许可的经验：旧许可还能用于新项目吗？指出适用限制、冲突和需要更新的内容。',as_of='2026-05-15T15:00:00+08:00')
            p.write_text(''.join(json.dumps(x,ensure_ascii=False)+'\n' for x in q))
        else:
            q=rows('04_queries/queries.jsonl');p.write_text('\n\n'.join(x['query_id']+' · '+x['as_of']+'\n'+x['prompt'] for x in q)+'\n')
    gold=ROOT/'05_private_gold_DO_NOT_INGEST/extraction_gold.jsonl';gs=[g for g in rows(str(gold.relative_to(ROOT))) if g['doc_id']!='F01-R2']
    for g in gs:g['expected_fields']['document_status']='unsigned_draft'
    revised=json.loads(json.dumps(next(g for g in gs if g['doc_id']=='F01')));revised['doc_id']='F01-R2';revised['expected_fields']['doc_id']='F01-R2';revised['source_clauses']=[x.replace('F01-','F01-R2-') for x in revised['source_clauses']]
    for fact in revised['key_obligations']:
        fact['clause_id']=fact['clause_id'].replace('F01-','F01-R2-')
        if fact['topic']=='delivery_delay':fact['value']='0.05%/day'
    gs.append(revised)
    gold.write_text(''.join(json.dumps(g,ensure_ascii=False)+'\n' for g in gs))
    recall=rows('05_private_gold_DO_NOT_INGEST/recall_gold.jsonl')
    for g in recall:
        if g['query_id']=='Q21':g.update(expected_decision='human_review',expected_evidence=['H03-C05','APPROVAL-H03-D1'],reasoning_requirements='真实读取合同及已知附条件审批记录；区分L1记录、未确认经验与正式经验。审批记录的存在不等于保函、生效和付款条件已满足；记忆确认不等于合同批准。')
        if g['query_id']=='Q22':g.update(required_memory_ids=['M06','M07'],expected_evidence=['WITHDRAWAL-H04-W1','APPROVAL-H04-D1'],reasoning_requirements='仅依据截至5月15日已知撤回及旧审批，识别旧先例不能授权新项目。观察旧正式经验、形成更新建议，并在后续人工编辑步骤核对版本。内部撤回不自动解除合同许可；不能引用尚未读入的六月政策。')
    (ROOT/'05_private_gold_DO_NOT_INGEST/recall_gold.jsonl').write_text(''.join(json.dumps(g,ensure_ascii=False)+'\n' for g in recall))
    dump('00_start_here/contract_manifest.json',manifest)
    dump('09_scenarios_OPERATOR_ONLY/source_locations.json',mappings)
    stages=json.loads((ROOT/'00_start_here/temporal_stages_OPERATOR_ONLY.json').read_text())
    for s in stages:
        s['isolation']='每个独立运行使用全新 Team 和 contract Agent；个人和团队 L2/L3 均按团队隔离。独立时点不可复用已读过未来事件的团队。'
        s['query_ids']=[q for q in s['query_ids'] if q not in ['Q21','Q22']]
        if s['stage']=='T1_before_guarantee':s['query_ids'].append('Q21')
        if s['stage']=='T4_internal_withdrawal_only':s['query_ids'].append('Q22')
    dump('00_start_here/temporal_stages_OPERATOR_ONLY.json',stages)
    inputs=json.loads((ROOT/'00_start_here/input_manifest_OPERATOR_ONLY.json').read_text());inputs=[x for x in inputs if x['input_id']!='F01-R2']
    inputs.append({'input_id':'F01-R2','kind':'revision_contract','available_at':'2026-08-21T09:00:00+08:00','path':revision['txt_path']})
    dump('00_start_here/input_manifest_OPERATOR_ONLY.json',inputs)
    print(json.dumps({'contracts':len(manifest),'events':len(events),'source_locations':len(mappings)},ensure_ascii=False))
if __name__=='__main__':main()
